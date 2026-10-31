import os
import sys
import json
import re
import urllib.request
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

USER_MAP = {
    "4mnesia": "Flavio Henriquez",
    "H4desDevs": "Flavio Henriquez",
    "Martin6k": "Martin Gutierrez",
    "martinkuruzg-hue": "Martin Gutierrez",
    "pedcampos": "Pedro Campos",
    "pedcamp": "Pedro Campos",
    "es-main77": "Esteban Maino"
}

def map_assignees(assignees):
    if not assignees:
        return "Sin asignar"
    mapped = []
    for a in assignees:
        login = a.get("login") if isinstance(a, dict) else a
        name = USER_MAP.get(login, login)
        if name not in mapped:
            mapped.append(name)
    if len(mapped) == 4 or set(mapped) == {"Flavio Henriquez", "Martin Gutierrez", "Pedro Campos", "Esteban Maino"}:
        return "Todo el equipo"
    return " / ".join(mapped)

def infer_tipo(title, labels):
    tl = title.lower()
    lbls = [l.get("name", "").lower() if isinstance(l, dict) else l.lower() for l in labels]
    if tl.startswith("épica") or tl.startswith("epica"):
        return "Épica"
    if tl.startswith("feat") or "enhancement" in lbls:
        return "Funcionalidad (Feat)"
    if tl.startswith("fix") or "bug" in lbls:
        return "Corrección (Fix)"
    if tl.startswith("chore"):
        return "Mantenimiento (Chore)"
    if tl.startswith("docs") or "documentation" in lbls:
        return "Documentación"
    if tl.startswith("security") or "seguridad" in tl:
        return "Seguridad"
    if tl.startswith("refactor"):
        return "Refactorización"
    if tl.startswith("ci") or "github_actions" in lbls:
        return "CI / DevOps"
    if "test" in tl:
        return "QA / Testing"
    if tl.startswith("decisión") or tl.startswith("decision") or "question" in lbls:
        return "Decisión Arquitectónica"
    if any(h in tl for h in ["manual", "diseño", "requisitos", "evidencias", "matriz"]):
        return "Hito / Entregable"
    return "Tarea Técnica"

def infer_modulo(title, labels):
    tl = title.lower()
    m = re.search(r'^[a-z_]+([^)]+)', tl)
    scope = m.group(1).strip() if m else ""
    parts = [p.strip() for p in scope.split(",") if p.strip()]
    mod_parts = []
    for p in parts:
        if p in ["portal", "frontend"]:
            mod_parts.append("Portal Web")
        elif p in ["admin"]:
            mod_parts.append("Portal Administrativo")
        elif p in ["ms-auth", "auth"]:
            mod_parts.append("ms-auth")
        elif p in ["ms-debt", "debt"]:
            mod_parts.append("ms-debt")
        elif p in ["ms-payments", "payments", "pasarelas"]:
            mod_parts.append("ms-payments")
        elif p in ["ms-ai", "ai"]:
            mod_parts.append("ms-ai")
        elif p in ["correo", "smtp"]:
            mod_parts.append("ms-auth / SMTP")
        elif p in ["docker"]:
            mod_parts.append("DevOps / Docker")
        elif p in ["config", "configuracion"]:
            mod_parts.append("DevOps / Configuración")
        elif p in ["base", "mysql"]:
            mod_parts.append("Base de Datos (MySQL)")
        elif p in ["readme"]:
            mod_parts.append("Documentación Técnica")
        elif p in ["integracion"]:
            mod_parts.append("Integración / Core")
        elif p in ["arquitectura"]:
            mod_parts.append("Arquitectura Backend")
        elif p in ["contrato", "legal"]:
            mod_parts.append("Gestión Proyecto APT")
        elif p in ["apofyx"]:
            mod_parts.append("ms-debt / APOFYX")
        elif p in ["patrimonio"]:
            mod_parts.append("Portal Web / Cartera")
        elif p in ["acceso"]:
            mod_parts.append("ms-auth")
        else:
            mod_parts.append(p)
            
    if mod_parts:
        res = []
        for mp in mod_parts:
            if mp not in res:
                res.append(mp)
        return " / ".join(res)
        
    if "portal" in tl:
        return "Portal Web"
    if "admin" in tl:
        return "Portal Administrativo"
    if "ms-auth" in tl:
        return "ms-auth"
    if "ms-debt" in tl:
        return "ms-debt"
    if "ms-payments" in tl or "pasarela" in tl or "pago" in tl:
        return "ms-payments"
    if "ms-ai" in tl or "ia" in tl or "asistente" in tl:
        return "ms-ai"
    if "docker" in tl:
        return "DevOps / Docker"
    if "mysql" in tl or "base" in tl:
        return "Base de Datos (MySQL)"
    if "vps" in tl or "cloud" in tl or "despliegue" in tl or "railway" in tl or "render" in tl:
        return "Cloud / DevOps"
    if any(k in tl for k in ["manual", "diseño", "requisitos", "evidencias", "matriz", "roles", "curso"]):
        return "Gestión Proyecto APT"
    return "Core Backend (Transversal)"

def fetch_issues(repo, token=None):
    url = f"https://api.github.com/repos/{repo}/issues?state=all&per_page=100"
    headers = {
        "Accept": "application/vnd.github.v3+json",
        "User-Agent": "DataBridge-Backlog-Sync"
    }
    if token:
        headers["Authorization"] = f"token {token}"
    req = urllib.request.Request(url, headers=headers)
    try:
        with urllib.request.urlopen(req) as resp:
            data = json.loads(resp.read().decode("utf-8"))
            return [i for i in data if "pull_request" not in i]
    except Exception as e:
        print(f"Error fetching {repo}: {e}")
        return []

def main():
    token = os.environ.get("GITHUB_TOKEN")
    repos = ["TechnicalBridge/TB_web", "TechnicalBridge/Technical-Bridge"]
    all_issues = []
    for r in repos:
        all_issues.extend(fetch_issues(r, token))
        
    print(f"Fetched {len(all_issues)} issues from GitHub.")
    if not all_issues:
        print("No issues found or error fetching.")
        return

    # Deduplicate by url
    unique_issues = {i['html_url']: i for i in all_issues}
    sorted_issues = sorted(unique_issues.values(), key=lambda x: (0 if x['state'] == 'open' else 1, -x['number']))

    # Load existing workbook for styling and user stories
    excel_path = os.path.join("docs", "Product Backlog - DataBridge.xlsx")
    wb = openpyxl.load_workbook(excel_path)
    ws_pb = wb["Product Backlog"]
    
    # Save headers
    headers = [ws_pb.cell(1, c).value for c in range(1, 11)]
    
    # Clear existing rows in Product Backlog
    while ws_pb.max_row > 1:
        ws_pb.delete_rows(2)

    # Styles
    font_data = Font(name="Arial", size=9)
    font_link = Font(name="Arial", size=9, color="2563EB", underline="single")
    fill_zebra_even = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_zebra_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border_side = Side(border_style="thin", color="E5E7EB")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    status_styles = {
        "Completado": {
            "fill": PatternFill(start_color="DEF7EC", end_color="DEF7EC", fill_type="solid"),
            "font": Font(name="Arial", size=9, bold=True, color="03543F")
        },
        "En curso": {
            "fill": PatternFill(start_color="E1EFFE", end_color="E1EFFE", fill_type="solid"),
            "font": Font(name="Arial", size=9, bold=True, color="1E429F")
        },
        "En revisión": {
            "fill": PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid"),
            "font": Font(name="Arial", size=9, bold=True, color="854D0E")
        },
        "Listo": {
            "fill": PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid"),
            "font": Font(name="Arial", size=9, bold=True, color="6B21A8")
        },
        "No iniciado": {
            "fill": PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid"),
            "font": Font(name="Arial", size=9, bold=True, color="4B5563")
        }
    }

    items = []
    for idx, iss in enumerate(sorted_issues, 1):
        pb_id = f"PB-{idx:02d}"
        tipo = infer_tipo(iss['title'], iss.get('labels', []))
        modulo = infer_modulo(iss['title'], iss.get('labels', []))
        titulo = iss['title'].strip()
        resp = map_assignees(iss.get('assignees', []))
        prio = "Alto" if ("security" in titulo.lower() or "épica" in titulo.lower()) else ("Bajo" if "chore" in titulo.lower() else "Medio")
        estado = "Completado" if iss['state'] == 'closed' else ("En curso" if resp != "Sin asignar" else "Listo")
        duracion = "XL" if tipo == "Épica" else ("M" if prio == "Alto" else "S")
        recurso = iss['html_url']
        obs = f"Issue #{iss['number']} de GitHub ({'Abierto' if iss['state'] == 'open' else 'Cerrado'})."

        row_vals = [pb_id, tipo, modulo, titulo, resp, prio, estado, duracion, recurso, obs]
        ws_pb.append(row_vals)
        r_idx = idx + 1
        ws_pb.row_dimensions[r_idx].height = 22
        base_fill = fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd
        
        for c_idx in range(1, 11):
            cell = ws_pb.cell(r_idx, c_idx)
            cell.font = font_data
            cell.fill = base_fill
            cell.border = thin_border
            if c_idx in [1, 2, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 9:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = font_link
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            if c_idx == 7 and estado in status_styles:
                cell.fill = status_styles[estado]["fill"]
                cell.font = status_styles[estado]["font"]

        items.append({
            'estado': estado,
            'prio': prio,
            'resp': resp,
            'modulo': modulo
        })

    # Save Excel
    wb.save(excel_path)
    print("Successfully updated Excel workbook.")

    # Also update CSV
    import csv
    csv_path = os.path.join("docs", "Product_Backlog_DataBridge.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for r in range(2, ws_pb.max_row + 1):
            writer.writerow([ws_pb.cell(r, c).value for c in range(1, 11)])
    print("Successfully updated CSV file.")

if __name__ == "__main__":
    main()
